import re
import pandas as pd # type: ignore
from openpyxl import load_workbook # type: ignore


def limpar_texto(valor):
    """
    Converte valores vazios/NaN para None
    e remove espaços desnecessários.
    """
    if pd.isna(valor):
        return None

    valor = str(valor).strip()

    if not valor:
        return None

    return valor



def extrair_urls_texto(valor):
    """Extrai todos os URLs de uma célula, preservando os links."""
    if pd.isna(valor):
        return []
    return [
        url.strip()
        for url in re.split(r"[\r\n]+", str(valor))
        if url.strip().lower().startswith(("http://", "https://"))
    ]


def extrair_hyperlinks_base(caminho):
    """Recupera hyperlinks da coluna Fotos da base original."""
    if not caminho:
        return []

    try:
        wb = load_workbook(caminho, read_only=False, data_only=False)
    except Exception as erro:
        raise RuntimeError(
            f"Falha ao abrir hyperlinks da base de fotos "
            f"'{caminho}': {erro}"
        ) from erro

    encontrados = []

    for ws in wb.worksheets:
        header_row = None
        col_matricula = None
        col_fotos = None

        for linha in range(1, min(ws.max_row, 20) + 1):
            valores = {
                coluna: (
                    str(ws.cell(linha, coluna).value).strip().lower()
                    if ws.cell(linha, coluna).value is not None
                    else ""
                )
                for coluna in range(1, ws.max_column + 1)
            }

            col_matricula = next(
                (c for c, v in valores.items() if v == "matrícula"),
                None
            )
            col_fotos = next(
                (c for c, v in valores.items() if v == "fotos"),
                None
            )

            if col_matricula is not None and col_fotos is not None:
                header_row = linha
                break

        if header_row is None:
            continue

        assert col_matricula is not None
        assert col_fotos is not None

        for linha in range(header_row + 1, ws.max_row + 1):
            matricula = limpar_texto(
                ws.cell(linha, col_matricula).value
            )
            hyperlink = ws.cell(linha, col_fotos).hyperlink

            if matricula and hyperlink and hyperlink.target:
                encontrados.append({
                    "Matrícula": matricula,
                    "Origem": "BASE_AG",
                    "Tipo": "GALERIA",
                    "Data do Registro": None,
                    "URL": hyperlink.target
                })

    wb.close()
    return encontrados


def construir_aba_fotos(
    base_atualizada,
    caminho_base_fotos=None
):
    """
    Cria a aba FOTOS com uma linha por matrícula.

    Cada foto fica em uma coluna:
    Matrícula | Foto 1 | Foto 2 | Foto 3 | ...

    Os links são preservados exatamente como vieram.
    """

    registros_fotos = []

    # =====================================================
    # 1. LINKS DA BASE ORIGINAL
    # =====================================================

    registros_fotos.extend(
        extrair_hyperlinks_base(
            caminho_base_fotos
        )
    )

    # =====================================================
    # 2. LINKS DO APPDEEP
    # =====================================================

    colunas_foto = [
        coluna
        for coluna in base_atualizada.columns
        if "foto" in str(coluna).lower()
    ]

    for _, linha in base_atualizada.iterrows():

        matricula = limpar_texto(
            linha.get("Matricula")
        )

        if not matricula:
            continue

        data_registro = linha.get(
            "Data do registro"
        )

        for coluna in colunas_foto:

            urls = extrair_urls_texto(
                linha.get(coluna)
            )

            for url in urls:

                registros_fotos.append({
                    "Matrícula": matricula,
                    "Origem": "APPDEEP",
                    "Tipo": str(coluna).strip(),
                    "Data do Registro": data_registro,
                    "URL": url
                })

    # =====================================================
    # 3. NENHUMA FOTO
    # =====================================================

    if not registros_fotos:

        return pd.DataFrame(
            columns=[
                "Matrícula"
            ]
        )

    fotos = pd.DataFrame(
        registros_fotos
    )

    # =====================================================
    # 4. LIMPEZA
    # =====================================================

    fotos["Matrícula"] = (
        fotos["Matrícula"]
        .map(limpar_texto)
    )

    fotos["URL"] = (
        fotos["URL"]
        .astype(str)
        .str.strip() # type: ignore
    )

    fotos = fotos[
        fotos["Matrícula"].notna()
        &
        fotos["URL"].str.startswith(
            ("http://", "https://"),
            na=False
        )
    ]

    # Remove links exatamente iguais
    # para a mesma matrícula
    fotos = fotos.drop_duplicates(
        subset=[
            "Matrícula",
            "URL"
        ]
    ) # type: ignore

    # =====================================================
    # 5. ORDENAÇÃO
    # =====================================================

    fotos = fotos.sort_values(
        by=[
            "Matrícula",
            "Data do Registro"
        ],
        na_position="last"
    ).reset_index(
        drop=True
    )

    # =====================================================
    # 6. NUMERAÇÃO DAS FOTOS
    # =====================================================

    fotos["NumeroFoto"] = (
        fotos
        .groupby("Matrícula")
        .cumcount()
        .add(1)
    )

    # =====================================================
    # 7. TRANSFORMA LINHAS EM COLUNAS
    # =====================================================

    fotos_wide = (
        fotos
        .pivot(
            index="Matrícula",
            columns="NumeroFoto",
            values="URL"
        )
        .reset_index()
    )

    # =====================================================
    # 8. RENOMEIA AS COLUNAS
    # =====================================================

    fotos_wide.columns = [
        (
            "Matrícula"
            if coluna == "Matrícula"
            else f"Foto {int(coluna)}"
        )
        for coluna in fotos_wide.columns
    ]

    # =====================================================
    # 9. ORDENA AS COLUNAS DE FOTOS
    # =====================================================

    colunas_foto_final = sorted(
        [
            coluna
            for coluna in fotos_wide.columns
            if coluna.startswith("Foto ")
        ],
        key=lambda coluna: int(
            coluna.replace("Foto ", "")
        )
    )

    fotos_wide = fotos_wide[
        ["Matrícula"] + colunas_foto_final
    ]

    return fotos_wide

def preparar_base(
    df,
    coluna_matricula,
    coluna_data,
    contar_visitas=False,
):
    """
    Prepara uma base para consolidação.

    Regras:
    1. Remove registros sem matrícula.
    2. Converte a data para datetime.
    3. Para a mesma matrícula:
       - data mais recente vence;
       - em empate, ADESÃO EXECUTADA vence;
       - novo empate, última ocorrência vence.
    """

    df = df.copy()

    # =====================================================
    # 1. MATRÍCULA
    # =====================================================

    df["_matricula"] = (
        df[coluna_matricula]
        .map(limpar_texto)
    )

    df = df[
        df["_matricula"].notna()
    ].copy()

    # =====================================================
    # 2. DATA
    # =====================================================

    df["_data_ordem"] = pd.to_datetime(
        df[coluna_data],
        errors="coerce",
        dayfirst=True
    )

    # =====================================================
    # 3. PRIORIDADE DE STATUS
    # =====================================================

    if "Status Adesão" in df.columns:

        df["_prioridade_status"] = (
            df["Status Adesão"]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.upper()
            .eq("ADESÃO EXECUTADA")
            .astype(int)
        )

    else:

        df["_prioridade_status"] = 0

    # =====================================================
    # 4. QUANTIDADE DE VISITAS
    # =====================================================

    if contar_visitas:
        # No APPDEEP, cada ocorrência da matrícula representa
        # uma visita. A contagem precisa ser feita ANTES
        # de remover os registros repetidos.
        df["_quantidade_visitas"] = (
            df.groupby("_matricula")["_matricula"]
            .transform("count")
        )
    elif "Qtd. Tentat." in df.columns:
        # Na BASE_AG, a quantidade oficial já vem pronta.
        df["_quantidade_visitas"] = pd.to_numeric(
            df["Qtd. Tentat."],
            errors="coerce"
        )
    else:
        df["_quantidade_visitas"] = None

    # =====================================================
    # 5. ORDEM ORIGINAL
    # =====================================================

    df["_ordem_original"] = range(len(df))

    # =====================================================
    # 5. ORDENAÇÃO
    # =====================================================

    df = df.sort_values(
        by=[
            "_matricula",
            "_data_ordem",
            "_prioridade_status",
            "_ordem_original"
        ],
        ascending=[
            True,
            True,
            True,
            True
        ],
        na_position="first"
    )

    # =====================================================
    # 6. MANTÉM O REGISTRO VENCEDOR
    # =====================================================

    df = (
        df
        .drop_duplicates(
            subset="_matricula",
            keep="last"
        )
        .set_index("_matricula")
    )

    return df



def link_registro(row):
    """Somente o link fornecido pela exportação, sem construir URLs por matrícula."""
    from urllib.parse import urlsplit
    if row is None:
        return None
    value = row.get("Link Backoffice")
    if value is None or pd.isna(value):
        return None
    value = str(value).strip()
    try:
        parsed = urlsplit(value)
        return value if parsed.scheme in {"https", "http"} and parsed.hostname else None
    except ValueError:
        return None

def consolidar(base_principal, base_atualizada, caminho_base_fotos=None):

    # =====================================================
    # 1. PREPARAÇÃO DAS BASES
    # =====================================================

    # BASE_AG.xlsx
    principal = preparar_base(
        base_principal,
        coluna_matricula="Matrícula",
        coluna_data="Data Adesão",
        contar_visitas=False
    )

    # APPDEEP
    atualizada = preparar_base(
        base_atualizada,
        coluna_matricula="Matricula",
        coluna_data="Data do registro",
        contar_visitas=True
    )

    # =====================================================
    # 2. MATRÍCULAS EXISTENTES NAS DUAS BASES
    # =====================================================

    matriculas = sorted(
        set(principal.index) |
        set(atualizada.index)
    )

    registros = []

    # =====================================================
    # 3. CONSOLIDAÇÃO
    # =====================================================

    for matricula in matriculas:

        p = (
            principal.loc[matricula]
            if matricula in principal.index
            else None
        )

        a = (
            atualizada.loc[matricula]
            if matricula in atualizada.index
            else None
        )

        # =================================================
        # 4. DETERMINA QUAL BASE É MAIS RECENTE
        # =================================================

        registro_vencedor = None
        registro_fallback = None

        if p is not None and a is not None:

            data_p = p.get("_data_ordem")
            data_a = a.get("_data_ordem")

            if pd.isna(data_p) and pd.isna(data_a):

                registro_vencedor = a
                registro_fallback = p

            elif pd.isna(data_p):

                registro_vencedor = a
                registro_fallback = p

            elif pd.isna(data_a):

                registro_vencedor = p
                registro_fallback = a

            elif data_a > data_p:

                registro_vencedor = a
                registro_fallback = p

            elif data_p > data_a:

                registro_vencedor = p
                registro_fallback = a

            else:

                # As duas bases não devem possuir a mesma data
                # de acordo com a regra definida para o ETL.
                registro_vencedor = a
                registro_fallback = p

        elif a is not None:

            registro_vencedor = a

        else:

            registro_vencedor = p

        # =================================================
        # 5. FUNÇÃO DE FALLBACK
        # =================================================

        def pegar_campo(coluna_vencedor, coluna_fallback=None):

            if registro_vencedor is not None:

                valor = limpar_texto(
                    registro_vencedor.get(coluna_vencedor)
                )

                if valor is not None:
                    return valor

            if (
                registro_fallback is not None
                and coluna_fallback is not None
            ):

                return limpar_texto(
                    registro_fallback.get(coluna_fallback)
                )

            return None

        # =================================================
        # 6. ENDEREÇO
        # =================================================

        if a is not None and registro_vencedor is a:

            endereco = limpar_texto(
                a.get("Endereço ")
            )

            if endereco is None and p is not None:
                endereco = limpar_texto(
                    p.get("Endereço")
                )

        else:

            endereco = (
                limpar_texto(
                    p.get("Endereço")
                )
                if p is not None
                else None
            )

            if endereco is None and a is not None:
                endereco = limpar_texto(
                    a.get("Endereço ")
                )

        # =================================================
        # 7. BAIRRO
        # =================================================

        # BASE_AG: Comunidade
        # APPDEEP: Bairro

        bairro = None

        if a is not None and registro_vencedor is a:
            bairro = limpar_texto(a.get("Bairro"))

            if bairro is None and p is not None:
                bairro = limpar_texto(p.get("Comunidade"))

        else:
            bairro = (
                limpar_texto(p.get("Comunidade"))
                if p is not None
                else None
            )

            if bairro is None and a is not None:
                bairro = limpar_texto(a.get("Bairro"))


        # =================================================
        # 8. STATUS ADESÃO
        # =================================================

        # BASE_AG: Status Adesão
        # APPDEEP: Situação Campo

        if a is not None and registro_vencedor is a:
            status_adesao = limpar_texto(
                a.get("Situação Campo")
            )

            if status_adesao is None and p is not None:
                status_adesao = limpar_texto(
                    p.get("Status Adesão")
                )
        else:
            status_adesao = (
                limpar_texto(p.get("Status Adesão"))
                if p is not None
                else None
            )

            if status_adesao is None and a is not None:
                status_adesao = limpar_texto(
                    a.get("Situação Campo")
                )

        # =================================================
        # 9. DATA ADESÃO

        # Rastreia a fonte efetiva do status, inclusive fallback por campo vazio.
        fonte_status = registro_vencedor
        coluna_status = "Situação Campo" if fonte_status is a else "Status Adesão"
        if fonte_status is not None and limpar_texto(fonte_status.get(coluna_status)) is None:
            coluna_fallback = "Status Adesão" if fonte_status is a else "Situação Campo"
            if registro_fallback is not None and limpar_texto(registro_fallback.get(coluna_fallback)) is not None:
                fonte_status = registro_fallback
        origem_adesao = ("APPDEEP" if fonte_status is a and a is not None else
                         "BASE_AG" if fonte_status is p and p is not None else None)
        data_origem = fonte_status.get("_data_ordem") if fonte_status is not None else pd.NaT

        # =================================================

        data_adesao = (
            p.get("Data Adesão")
            if p is not None
            else None
        )

        # =================================================
        # 10. SITUAÇÃO BACKOFFICE
        # =================================================

        situacao_backoffice = (
            limpar_texto(
                a.get("Situação Backoffice")
            )
            if a is not None
            else None
        )

        # =================================================
        # 11. DATA DO REGISTRO
        # =================================================

        data_registro = (
            a.get("Data do registro")
            if a is not None
            else None
        )

        # =================================================
        # 12. EQUIPE EXEC.
        # =================================================

        equipe = (
            limpar_texto(
                p.get("Equipe Exec.")
            )
            if p is not None
            else None
        )

        # =================================================
        # 13. AGENTE
        # =================================================

        agente = (
            limpar_texto(
                a.get("Nome do agente ")
            )
            if a is not None
            else None
        )

        # =================================================
        # 14. LATITUDE
        # =================================================

        latitude = None

        if a is not None and registro_vencedor is a:

            latitude = a.get("Latitude")

            if pd.isna(latitude) and p is not None:
                latitude = p.get("Latitude")

        else:

            latitude = (
                p.get("Latitude")
                if p is not None
                else None
            )

            if pd.isna(latitude) and a is not None:
                latitude = a.get("Latitude")

        # =================================================
        # 15. LONGITUDE
        # =================================================

        longitude = None

        if a is not None and registro_vencedor is a:

            longitude = a.get("Longitude")

            if pd.isna(longitude) and p is not None:
                longitude = p.get("Longitude")

        else:

            longitude = (
                p.get("Longitude")
                if p is not None
                else None
            )

            if pd.isna(longitude) and a is not None:
                longitude = a.get("Longitude")

        # =================================================
        # 16. NOME CLIENTE
        # =================================================

        nome_cliente = None

        if a is not None and registro_vencedor is a:

            nome_cliente = limpar_texto(
                a.get("Nome Cliente")
            )

            if nome_cliente is None and p is not None:
                nome_cliente = limpar_texto(
                    p.get("Nome Cliente")
                )

        else:

            nome_cliente = (
                limpar_texto(
                    p.get("Nome Cliente")
                )
                if p is not None
                else None
            )

            if nome_cliente is None and a is not None:
                nome_cliente = limpar_texto(
                    a.get("Nome Cliente")
                )

        # =================================================
        # 17. TELEFONE
        # =================================================

        telefone = None

        if a is not None and registro_vencedor is a:

            telefone = limpar_texto(
                a.get("Telefone")
            )

            if telefone is None and p is not None:
                telefone = limpar_texto(
                    p.get("Telefone")
                )

        else:

            telefone = (
                limpar_texto(
                    p.get("Telefone")
                )
                if p is not None
                else None
            )

            if telefone is None and a is not None:
                telefone = limpar_texto(
                    a.get("Telefone")
                )

        # =================================================
        # 18. QUANTIDADE DE VISITAS
        # =================================================

        # REGRA:
        # - Se existe na BASE_AG, usa "Qtd. Tentat.".
        # - Se existe somente no APPDEEP, usa a quantidade
        #   de vezes que a matrícula apareceu no APPDEEP.
        #
        # A BASE_AG tem prioridade quando a matrícula existe
        # nas duas bases.

        if p is not None:
            quantidade_visitas = p.get("_quantidade_visitas")
        elif a is not None:
            quantidade_visitas = a.get("_quantidade_visitas")
        else:
            quantidade_visitas = None

        if pd.isna(quantidade_visitas):
            quantidade_visitas = None



        # =================================================
        # 19. ORIGEM
        # =================================================

        if p is not None and a is not None:

            origem = "AMBOS"

        elif p is not None:

            origem = "BASE PRINCIPAL"

        else:

            origem = "BASE ATUALIZADA"

        # =================================================
        # 20. REGISTRO FINAL
        # =================================================

        registros.append({
            "Matrícula": matricula,
            "Endereço": endereco,
            "Bairro": bairro,
            "Status Adesão": status_adesao,
            "Data Adesão": data_adesao,
            "Situação Backoffice": situacao_backoffice,
            "Data do Registro": data_registro,
            "Equipe Exec.": equipe,
            "Agente": agente,
            "Latitude": latitude,
            "Longitude": longitude,
            "Nome Cliente": nome_cliente,
            "Telefone": telefone,
            "Quantidade de Visitas": quantidade_visitas,
            "origem_adesao": origem_adesao,
            "data_origem_adesao": data_origem,
            "link_appdeep": link_registro(fonte_status) if origem_adesao == "APPDEEP" else None,
            "registro_appdeep": limpar_texto(a.get("Code Deep")) if origem_adesao == "APPDEEP" else None,
            "Origem dos Dados": origem
        })

    # =====================================================
    # 21. DATAFRAME FINAL
    # =====================================================

    resultado = pd.DataFrame(registros)

    # =====================================================
    # 22. PADRONIZAÇÃO FINAL DOS TEXTOS
    # =====================================================

    def padronizar_saida(valor):
        if pd.isna(valor):
            return None

        if not isinstance(valor, str):
            return valor

        # Remove espaços extras e coloca em maiúsculo.
        valor = " ".join(valor.strip().split())

        return valor.upper() if valor else None

    # Padronizamos somente os campos de texto definidos
    # para a BASE_AGUA. Não aplicamos isso na aba FOTOS,
    # principalmente nos URLs.
    colunas_texto = [
        "Matrícula",
        "Endereço",
        "Bairro",
        "Status Adesão",
        "Situação Backoffice",
        "Equipe Exec.",
        "Agente",
        "Nome Cliente",
        "Origem dos Dados"
    ]

    for coluna in colunas_texto:
        if coluna in resultado.columns:
            resultado[coluna] = resultado[coluna].map(
                padronizar_saida
            )

    # =====================================================
    # 23. ABA FOTOS
    # =====================================================

    fotos = construir_aba_fotos(
        base_atualizada,
        caminho_base_fotos
    )

    return resultado, fotos
